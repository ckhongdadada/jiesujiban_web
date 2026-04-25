import docx
import sys

def read_docx(file_path):
    doc = docx.Document(file_path)
    output_path = "docx_content.txt"
    with open(output_path, "w", encoding="utf-8") as f:
        for para in doc.paragraphs:
            f.write(para.text + "\n")
        f.write("\n\n--- TABLES ---\n")
        for table in doc.tables:
            for row in table.rows:
                f.write("\t".join([cell.text for cell in row.cells]) + "\n")
            f.write("\n")
    print(f"Written to {output_path}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        read_docx(sys.argv[1])
    else:
        print("Please provide a file path.")
