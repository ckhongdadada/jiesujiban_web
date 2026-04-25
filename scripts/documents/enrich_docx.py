import docx
from docx.shared import Pt
import sys

def enrich_document(file_path):
    doc = docx.Document(file_path)
    
    # We will append the new content before the appendices or at the end.
    # Actually, appending at the end is safest. Let's just append new sections and then move them before Appendix if possible.
    # It's easier to just add them at the end of the main body, let's say after "10. 测试、指标与解释" or before "11. 当前系统的优势与边界".
    
    # Let's find the paragraph index for "11. 当前系统的优势与边界"
    insert_idx = len(doc.paragraphs)
    for i, p in enumerate(doc.paragraphs):
        if "11. 当前系统的优势与边界" in p.text:
            insert_idx = i
            break
            
    # We can't insert easily with python-docx (it supports insert_paragraph_before on a paragraph object)
    if insert_idx < len(doc.paragraphs):
        target_p = doc.paragraphs[insert_idx]
    else:
        target_p = None
        
    def add_heading(text, level=2, target=None):
        if target:
            return target.insert_paragraph_before(text, style=f'Heading {level}')
        else:
            return doc.add_heading(text, level=level)
            
    def add_paragraph(text, target=None):
        if target:
            return target.insert_paragraph_before(text)
        else:
            return doc.add_paragraph(text)

    # 10.1 混合检索 (Hybrid Retrieval)
    add_heading("10.1 混合检索与知识图谱融合 (Hybrid Retrieval & Knowledge Graph)", level=2, target=target_p)
    add_paragraph("为了突破单一稠密向量检索（Dense Retrieval）在专有名词和政务数字编号上的匹配瓶颈，系统在原有BGE模型基础上引入了混合检索机制（Hybrid Retriever, 参见 enhancements/rag_hybrid_retriever.py）。通过结合TF-IDF/BM25的稀疏检索能力与大语言模型的结构化知识库（Structured KB & Knowledge Graph），实现了“精准关键词拦截+泛化语义召回+图谱关系推理”的三路召回架构。这有效改善了特定政策条文和罕见地名的召回率。", target=target_p)
    
    # 10.2 事实验证 (Fact Extraction and Verification)
    add_heading("10.2 严格的事实验证机制 (Fact Extraction & Verification)", level=2, target=target_p)
    add_paragraph("在政务场景中大模型幻觉是零容忍的。系统新增了独立的事实验证模块（fact_extractor.py, fact_verifier.py），采用“生成后校验”的策略。针对大模型生成的回复初稿，提取其中的核心“事实元组”（如处理状态、时间节点、政策指标等），然后与召回的RAG参考文档及知识图谱进行硬链接对比验证。一旦发现虚假套话或“捏造现场核查结果”，引擎将自动打回重写或给用户提出高风险红色预警。", target=target_p)

    # 10.3 推测解码与性能提速 (Speculative Decoding & Async/Cache)
    add_heading("10.3 并发优化与推测解码提速 (Performance & Speculative Decoding)", level=2, target=target_p)
    add_paragraph("在工程性能方面，平台通过引入异步任务处理器（async_processor.py）和多级缓存管理器（cache_manager.py），实现了高并发下的毫秒级响应能力，拦截大量重复或热点投诉问题。针对Qwen大语言模型的生成瓶颈，系统还实验性地集成了推测解码机制（Speculative Decoding, 参见 speculative_decoder.py），利用小模型快速草拟Token序列并交由大模型并行验证，使吞吐量和生成速度获得了显著提升。", target=target_p)

    # 10.4 安全与数据隐私 (Security & Privacy)
    add_heading("10.4 接口安全与敏感数据隐私保护 (Security & Privacy Protection)", level=2, target=target_p)
    add_paragraph("考虑到系统处理的往往是包含个人隐私或者敏感地点的实名投诉信息，系统新增强化了安全模块（security/auth.py, security/encryption.py）。这包括基于令牌认证的严格API访问控制（RBAC与JWT机制集成）以及落库数据的脱敏与非对称加密处理，保证端到端政务数据流转的合规与防泄露要求。", target=target_p)

    doc.save(file_path)
    print(f"Document {file_path} enriched successfully.")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        enrich_document(sys.argv[1])
    else:
        print("Please provide a file path.")
